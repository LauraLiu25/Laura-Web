from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "assets" / "documents"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "个人网站PRD 3.0.docx"


FONT_CN = "Microsoft YaHei"
FONT_SERIF = "SimSun"
INK = RGBColor(30, 30, 30)
MUTED = RGBColor(92, 92, 92)
HEADER_FILL = "243447"
LIGHT_FILL = "F4F6F8"
PALE_FILL = "F8FAFC"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=None, bold=None, color=None, name=FONT_CN):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_para(doc, text="", style=None, size=10.5, bold=False, color=INK, space_after=6, align=None):
    p = doc.add_paragraph(style=style)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.18
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        set_font(r, size=15 if level == 1 else 12, bold=True, color=INK, name=FONT_CN)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55 + level * 0.45)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.14
    r = p.add_run(text)
    set_font(r, size=10.2, color=INK)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr[i]
        set_cell_shading(cell, HEADER_FILL)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        set_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
        if widths:
            cell.width = Cm(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            set_cell_shading(cell, PALE_FILL if row_index % 2 else "FFFFFF")
            set_cell_borders(cell)
            set_cell_margins(cell, top=110, start=120, bottom=110, end=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(text))
            set_font(run, size=9.3, color=INK)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(2.05)
    section.right_margin = Cm(2.05)

    styles = doc.styles
    styles["Normal"].font.name = FONT_CN
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("个人网站产品需求文档 3.0")
    set_font(r, size=20, bold=True, color=INK)

    subtitle = add_para(
        doc,
        "面向电商运营平台方与品牌方岗位投递的个人作品型网站",
        size=10.5,
        color=MUTED,
        space_after=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    meta_rows = [
        ["产品名称", "刘红锁 Laura 个人网站"],
        ["当前版本", "3.0"],
        ["目标场景", "电商运营、商品运营、平台运营、品牌运营相关岗位投递"],
        ["核心定位", "在商品、数据与用户之间，寻找更稳健的运营增长路径"],
        ["文档用途", "梳理网站目标、结构、核心功能、内容策略与后续迭代方向"],
    ]
    add_table(doc, ["字段", "说明"], meta_rows, widths=[3.2, 12.5])

    add_heading(doc, "一 项目背景", 1)
    add_para(
        doc,
        "3.0 版本的个人网站从通用型个人展示，调整为更服务于求职投递的作品型网站。网站重点面向电商运营平台方与品牌方相关岗位，核心任务是让招聘方在较短时间内理解候选人的运营实践、数据意识、商业判断和 AI 工具应用能力。",
        space_after=6,
    )
    add_para(
        doc,
        "与 1.0 版本相比，3.0 不再只承担简历展示功能，而是补充首页岗位匹配入口、留言板转化入口、东方优选电商运营经历、项目正式发布链接与更清晰的模块化表达。整体目标是将“我做过什么”升级为“这些经历如何对应目标岗位能力”。",
        space_after=6,
    )

    add_heading(doc, "二 版本演进", 1)
    version_rows = [
        ["1.0", "基础个人网站", "完成首页、关于我、教育经历、工作经历、项目经历等基础展示；尚未包含首页岗位匹配与留言板。"],
        ["2.0", "求职投递增强", "增加 AI 岗位匹配、站点内容重构、项目经历展示优化，开始将网站作为求职作品集使用。"],
        ["3.0", "电商运营定位强化", "围绕电商运营岗位重写首页与关于我，新增东方优选经历页，工作经历采用卡片与弹窗，项目链接更新为正式发布入口。"],
    ]
    add_table(doc, ["版本", "阶段目标", "主要变化"], version_rows, widths=[2.0, 3.6, 10.2])

    add_heading(doc, "三 用户与使用场景", 1)
    persona_rows = [
        ["招聘方 HR", "快速判断候选人是否匹配岗位", "首页定位、AI 岗位匹配器、简历下载、联系方式"],
        ["业务面试官", "查看真实经历与可迁移能力", "工作经历、项目经历、弹窗详情、量化成果"],
        ["候选人本人", "统一管理投递展示材料", "模块化内容、项目链接、留言板、线上作品入口"],
    ]
    add_table(doc, ["用户", "核心需求", "对应功能"], persona_rows, widths=[3.0, 5.0, 7.8])

    add_heading(doc, "四 产品目标", 1)
    add_bullet(doc, "建立清晰的电商运营求职定位，突出商品、数据、用户与流程意识。")
    add_bullet(doc, "把经历从简历式堆叠转为模块化证据，提升招聘方浏览效率。")
    add_bullet(doc, "通过 AI 岗位匹配器与留言板形成从访问、判断到联系的闭环。")
    add_bullet(doc, "用项目经历补充运营之外的商业分析、AI 应用和结构化表达能力。")

    add_heading(doc, "五 信息架构", 1)
    ia_rows = [
        ["首页", "建立第一印象，提供定位、简历下载、岗位匹配入口", "访问后 5 秒内理解候选人方向"],
        ["关于我", "围绕商品运营、经营数据、用户转化重写个人介绍", "让个人定位贴合电商运营岗位"],
        ["教育经历", "展示本科、支教、硕士经历与荣誉", "补充学习能力、责任感和长期投入"],
        ["工作经历", "呈现东方优选电商运营实习与天职审计实习", "突出商品运营、数据分析、流程风控与财务基础"],
        ["项目经历", "展示个人网站、极因造物、宇树科技案例", "补充作品能力、商业分析和正式成果"],
        ["更多", "联系方式、Github、小红书、留言板", "完成联系与反馈转化"],
    ]
    add_table(doc, ["模块", "内容", "目标"], ia_rows, widths=[2.2, 7.2, 6.3])

    add_heading(doc, "六 核心功能需求", 1)
    feature_rows = [
        ["F1 首页定位", "展示姓名、核心 slogan、简历预览与 AI 岗位匹配入口。", "高"],
        ["F2 AI 岗位匹配器", "输入 JD 后输出候选人与岗位匹配分析，帮助 HR 快速理解适配点。", "高"],
        ["F3 关于我能力卡", "将能力从“产品与效率”调整为更贴合电商运营的表达。", "高"],
        ["F4 工作经历卡片", "东方优选与审计经历均采用三卡结构，标题、编号、间距和视觉风格统一。", "高"],
        ["F5 工作详情弹窗", "点击东方优选卡片后，以短句、指标、店铺标签和职责成果模块展示详情。", "高"],
        ["F6 项目详情弹窗", "个人网站展示六大模块，宇树案例点击后跳转官方发布页，极因造物保持现有展示。", "中"],
        ["F7 留言板", "提供访客留言入口，补充联系方式之外的互动路径。", "中"],
    ]
    add_table(doc, ["编号", "需求说明", "优先级"], feature_rows, widths=[1.7, 11.4, 2.2])

    add_heading(doc, "七 3.0 内容策略", 1)
    add_heading(doc, "首页与关于我", 2)
    add_bullet(doc, "首页 slogan 使用“在商品、数据与用户之间，寻找更稳健的运营增长路径”。")
    add_bullet(doc, "关于我聚焦商品运营、经营数据与用户转化，弱化与当前投递方向关联较低的产品效率叙述。")
    add_bullet(doc, "能力卡片突出电商运营与增长、数据分析、AI 应用和跨团队协同。")

    add_heading(doc, "工作经历", 2)
    work_rows = [
        ["东方优选", "2026年7月-2026年9月", "电商运营实习生", "商品全链路运营、AI 提效与价格风控、库存与经营数据优化"],
        ["天职国际", "2025年12月-2026年3月", "审计实习生", "内控测试、审计底稿编制、集团合并底稿编制"],
    ]
    add_table(doc, ["公司", "时间", "岗位", "展示重点"], work_rows, widths=[3.0, 3.8, 3.0, 5.8])
    add_bullet(doc, "东方优选详情中将 6 家店铺拆成标签：天猫旗舰店、淘宝主店、天猫生鲜、天猫保健、淘宝家居日用、淘宝农场。")
    add_bullet(doc, "东方优选第二张卡片加入千牛上传助手示意图，表达自动新增素材文件夹、识别主图/详情页/后台图并分别上传的工具能力。")
    add_bullet(doc, "两段工作经历的三张卡片统一标题、编号、间距与图标风格。")

    add_heading(doc, "项目经历", 2)
    add_bullet(doc, "个人网站项目只保留个人网站，不再展示辅助求职平台，配图转为六大模块示意。")
    add_bullet(doc, "宇树科技案例已正式发布，详情入口直接跳转中国管理案例共享中心官网发布页。")
    add_bullet(doc, "极因造物 Aegle X 保持现有内容，继续承担商业模式设计与国家级竞赛成果展示。")

    add_heading(doc, "八 关键交互流程", 1)
    flow_rows = [
        ["招聘方进入首页", "浏览定位与 slogan，选择下载简历或输入 JD 进行匹配", "明确候选人方向并进入下一步判断"],
        ["查看工作经历", "浏览东方优选/审计三张卡片，点击卡片查看弹窗详情", "用短时间获得真实任务和量化成果"],
        ["查看项目经历", "点击项目卡片，查看网站架构、项目成果或官方链接", "补充候选人的结构化表达和商业分析能力"],
        ["产生联系意向", "通过联系方式复制、Github/小红书跳转或留言板提交", "完成访问到联系的转化"],
    ]
    add_table(doc, ["流程", "用户动作", "产品目标"], flow_rows, widths=[3.6, 6.4, 5.8])

    add_heading(doc, "九 非功能需求", 1)
    add_bullet(doc, "响应式展示：桌面端优先保证一屏浏览效率，移动端避免卡片过长和文字拥挤。")
    add_bullet(doc, "视觉一致性：工作经历中的编号、小标题、卡片间距、图标样式保持一致。")
    add_bullet(doc, "内容可维护：图片、项目链接、岗位信息和经历指标应可独立替换，减少后续修改成本。")
    add_bullet(doc, "加载与可访问性：图片需设置替代文本，弹窗支持关闭按钮、遮罩关闭和 Esc 关闭。")

    add_heading(doc, "十 验收标准", 1)
    accept_rows = [
        ["首页定位", "slogan 与关于我内容均指向电商运营相关岗位", "通过"],
        ["工作经历", "东方优选位于审计前一页，两段经历日期和岗位格式统一", "通过"],
        ["东方优选卡片", "三张卡片有配图，标题不换行，关键词位于图片上方，点击弹窗展示短块内容", "通过"],
        ["审计卡片", "三张卡片与东方优选在标题、编号和间距风格上保持一致", "通过"],
        ["项目经历", "个人网站项目只保留个人网站，宇树案例跳转官方发布页", "通过"],
        ["转化入口", "保留简历预览、AI 岗位匹配器、联系方式和留言板", "通过"],
    ]
    add_table(doc, ["验收项", "标准", "状态"], accept_rows, widths=[3.4, 9.6, 2.0])

    add_heading(doc, "十一 后续迭代建议", 1)
    add_bullet(doc, "补充 3 张工作经历配图的最终版本，并统一命名到 work 图片目录，降低后续维护成本。")
    add_bullet(doc, "将 AI 岗位匹配器输出进一步压缩为 HR 友好的“匹配结论、证据、风险、建议提问”。")
    add_bullet(doc, "为不同岗位方向准备轻量化内容开关，例如电商运营版、平台运营版、商业分析版。")
    add_bullet(doc, "增加上线后的访问数据复盘，用点击率和留言转化判断首页与项目卡片是否有效。")

    add_para(
        doc,
        "本 PRD 对应 3.0 版本当前实现状态，后续可根据实际投递岗位和访问反馈继续调整页面优先级、内容措辞与项目呈现方式。",
        size=9.5,
        color=MUTED,
        space_after=0,
    )

    doc.core_properties.title = "个人网站产品需求文档 3.0"
    doc.core_properties.author = "刘红锁"
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build_doc())
